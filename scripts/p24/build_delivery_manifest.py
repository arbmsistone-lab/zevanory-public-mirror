#!/usr/bin/env python3
"""Gera data/p24/delivery-manifest.json a partir dos zips dos produtos (execução local, fora do CI).

Os zips NÃO são versionados. Este script:
  1. lê o catálogo do Worker (sku, artifact_name, artifact_sha256) e localiza cada zip em --zips;
  2. confere o SHA-256 do zip contra o catálogo (divergência aborta);
  3. testa a integridade do zip e abre cada arquivo em memória:
       PDF  -> pypdf: páginas, texto, marcação (tagged), título, idioma, criptografia;
       DOCX -> python-docx: parágrafos, tabelas, título;
       XLSX -> openpyxl: abas, fórmulas; recálculo pelo LibreOffice (se instalado) e contagem de erros;
       TXT  -> UTF-8 estrito, sem caractere de substituição;
  4. conta, por arquivo, as âncoras temáticas citadas em data/p24/promise-map.json;
  5. procura marcadores de rascunho e promessas de ganho (achado para revisão, nunca correção automática).

Grava só metadados (caminho, tamanho, SHA-256, resultado da abertura, contagens). Nenhum texto integral do
material é copiado para o repositório.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("p24check", HERE / "product_quality_check.py")
p24 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(p24)  # type: ignore[union-attr]

# Maiúsculas exatas: "todo" é palavra comum em português e não marca rascunho.
PLACEHOLDER = re.compile(r"\b(TODO|FIXME|XXX|TBD)\b|(?i:lorem ipsum)|\{\{|\}\}|(?i:\[inserir)")
INCOME = re.compile(r"(lucro|renda|resultado|faturamento|venda)s? garantid|fique rico|ganhe dinheiro|renda extra|ganhe r\$", re.I)
XLSX_ERRORS = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")


def sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def kind_of(path: str) -> str:
    ext = path.rsplit(".", 1)[-1].lower() if "." in path else ""
    return {"pdf": "pdf", "docx": "docx", "xlsx": "xlsx", "txt": "txt"}.get(ext, "other")


def open_pdf(data: bytes) -> tuple[dict, str]:
    import pypdf
    r = pypdf.PdfReader(io.BytesIO(data))
    if r.is_encrypted:
        return {"ok": False, "detail": "PDF criptografado"}, ""
    text = "\n".join((pg.extract_text() or "") for pg in r.pages)
    root = r.trailer["/Root"]
    mark = root.get("/MarkInfo")
    tagged = bool(mark and mark.get_object().get("/Marked")) and "/StructTreeRoot" in root
    meta = r.metadata or {}
    title = str(meta.get("/Title") or "").strip()
    lang = str(root.get("/Lang") or "").strip()
    ok = len(r.pages) > 0 and len(text.strip()) > 0
    return {"ok": ok, "detail": f"PDF {len(r.pages)} página(s), {len(text)} caracteres de texto extraído", "pages": len(r.pages),
            "tagged": tagged, "title": title or None, "lang": lang or None}, text


def open_docx(data: bytes) -> tuple[dict, str]:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    text = "\n".join(parts)
    title = (d.core_properties.title or "").strip()
    lang = None
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if "word/styles.xml" in z.namelist():
            styles = z.read("word/styles.xml").decode("utf-8", "replace")
            m = re.search(r"<w:docDefaults>.*?<w:lang\b[^>]*\bw:val=\"([^\"]+)\"", styles, re.S)
            lang = m.group(1) if m else None
    ok = len(text.strip()) > 0
    return {"ok": ok, "detail": f"DOCX {len(d.paragraphs)} parágrafo(s), {len(d.tables)} tabela(s)", "paragraphs": len(d.paragraphs),
            "tables": len(d.tables), "title": title or None, "lang": lang}, text


def recalc_xlsx(data: bytes, name: str, workdir: Path) -> dict:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {"status": "NOT_AVAILABLE", "detail": "LibreOffice ausente"}
    import openpyxl
    src = workdir / "in" / Path(name).name
    out = workdir / "out"
    src.parent.mkdir(parents=True, exist_ok=True)
    out.mkdir(parents=True, exist_ok=True)
    src.write_bytes(data)
    prof = workdir / "lo-profile"
    cmd = [soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)]
    try:
        subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    except (subprocess.SubprocessError, OSError) as exc:
        return {"status": "FAIL", "detail": f"LibreOffice não converteu: {exc.__class__.__name__}", "errors": 1}
    wb = openpyxl.load_workbook(out / src.name, data_only=True)
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip().startswith(XLSX_ERRORS):
                    errs.append(f"{ws.title}!{c.coordinate}={c.value.strip()}")
    return {"status": "PASS" if not errs else "FAIL", "engine": "LibreOffice headless", "errors": len(errs), "examples": errs[:10]}


def open_xlsx(data: bytes, name: str, workdir: Path) -> tuple[dict, str]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    formulas, literal_err, strings = 0, 0, list(wb.sheetnames)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    if v.startswith("="):
                        formulas += 1
                    else:
                        strings.append(v)
                        if v.strip().startswith(XLSX_ERRORS):
                            literal_err += 1
    rc = recalc_xlsx(data, name, workdir)
    ok = len(wb.sheetnames) > 0 and literal_err == 0
    return {"ok": ok, "detail": f"XLSX {len(wb.sheetnames)} aba(s), {formulas} fórmula(s)", "sheets": wb.sheetnames, "formulas": formulas,
            "literal_errors": literal_err, "recalc": rc}, "\n".join(strings)


def open_txt(data: bytes) -> tuple[dict, str]:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        return {"ok": False, "detail": f"não é UTF-8: {exc.reason}"}, ""
    ok = bool(text.strip()) and "�" not in text
    return {"ok": ok, "detail": f"TXT UTF-8, {len(text.splitlines())} linha(s)"}, text


def quality_hits(text: str) -> list[str]:
    hits = []
    for rx, label in ((PLACEHOLDER, "marcador de rascunho"), (INCOME, "possível promessa de ganho")):
        for m in rx.finditer(text):
            a, b = max(0, m.start() - 50), min(len(text), m.end() + 50)
            hits.append(f"{label}: …{re.sub(r'\s+', ' ', text[a:b])}…")
    return hits[:10]


def anchors_by_file(pmap: dict) -> dict[str, dict[str, set[str]]]:
    out: dict[str, dict[str, set[str]]] = {}
    for slug, prod in pmap["products"].items():
        for p in prod["promises"]:
            for ev in p.get("evidence", []) or []:
                out.setdefault(slug, {}).setdefault(ev["file"], set()).update(p24.norm_anchor(a) for a in ev["anchors"])
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--root", default=".")
    ap.add_argument("--zips", required=True, help="pasta com os zips dos produtos")
    ap.add_argument("--out", default="data/p24/delivery-manifest.json")
    ap.add_argument("--generated-at", required=True, help="data/hora ISO da conferência")
    ap.add_argument("--source-label", default="cópia local dos zips catalogados")
    args = ap.parse_args(argv)
    root, zips = Path(args.root), Path(args.zips)
    pmap = json.loads((root / "data/p24/promise-map.json").read_text(encoding="utf-8"))
    catalog = p24.parse_worker_catalog((root / "worker/cloudflare-worker.recovered.mjs").read_text(encoding="utf-8"))
    anchors = anchors_by_file(pmap)
    import pypdf
    products = {}
    with tempfile.TemporaryDirectory(prefix="p24-") as td:
        for slug, prod in pmap["products"].items():
            cat = catalog[prod["sku"]]
            zpath = zips / cat["artifact_name"]
            blob = zpath.read_bytes()
            digest = sha256(blob)
            if digest != cat["artifact_sha256"]:
                print(f"ABORT {slug}: SHA-256 {digest} difere do catálogo {cat['artifact_sha256']}")
                return 2
            zf = zipfile.ZipFile(io.BytesIO(blob))
            bad = zf.testzip()
            entries = []
            for info in sorted(zf.infolist(), key=lambda i: i.filename):
                if info.is_dir():
                    continue
                data = zf.read(info.filename)
                k = kind_of(info.filename)
                try:
                    if k == "pdf":
                        o, text = open_pdf(data)
                    elif k == "docx":
                        o, text = open_docx(data)
                    elif k == "xlsx":
                        o, text = open_xlsx(data, info.filename, Path(td) / slug / str(len(entries)))
                    elif k == "txt":
                        o, text = open_txt(data)
                    else:
                        o, text = {"ok": False, "detail": "formato não previsto na rubrica"}, ""
                except Exception as exc:  # qualquer falha de abertura é resultado, não exceção
                    o, text = {"ok": False, "detail": f"falha ao abrir: {exc.__class__.__name__}"}, ""
                ntext = p24.norm_anchor(text)
                want = sorted(anchors.get(slug, {}).get(info.filename, set()))
                entries.append({
                    "path": info.filename,
                    "size": len(data),
                    "sha256": sha256(data),
                    "kind": k,
                    "open": o,
                    "anchors": {a: ntext.count(a) for a in want},
                    "quality_hits": quality_hits(text),
                })
            products[slug] = {
                "sku": prod["sku"],
                "artifact_name": cat["artifact_name"],
                "artifact_sha256": digest,
                "artifact_bytes": len(blob),
                "zip_integrity": "ok" if bad is None else f"corrompido: {bad}",
                "entries": entries,
            }
            print(f"{slug}: {len(entries)} arquivo(s), sha ok, zip {'ok' if bad is None else 'CORROMPIDO'}")
    out = {
        "schema": p24.SCHEMA_MANIFEST,
        "generated_at": args.generated_at,
        "source": args.source_label,
        "generator": "scripts/p24/build_delivery_manifest.py",
        "extractors": {"pdf": f"pypdf {pypdf.__version__}", "docx": "python-docx", "xlsx": "openpyxl + LibreOffice headless (recálculo)", "txt": "utf-8 estrito"},
        "note": "Zips não versionados. Este manifesto guarda só metadados; o SHA-256 de cada zip é conferido contra o catálogo do Worker no CI.",
        "products": products,
    }
    Path(root / args.out).write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
